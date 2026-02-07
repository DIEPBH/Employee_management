# services/excel_utils.py
from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple, Any
import openpyxl


@dataclass
class ExcelSheetData:
    header: List[str]               # header normalized
    col: Dict[str, int]             # column name -> 0-based index
    rows: List[Tuple[Any, ...]]     # values_only rows (from row 2)


def normalize_header(v) -> str:
    return (str(v).strip().lower() if v is not None else "")


def load_excel_sheet(file_obj, sheet_name: Optional[str] = None) -> ExcelSheetData:
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active

    header = [normalize_header(c.value) for c in ws[1]]
    col = {name: idx for idx, name in enumerate(header) if name}

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return ExcelSheetData(header=header, col=col, rows=rows)


def validate_required_columns(col_map: Dict[str, int], required: List[str]) -> List[str]:
    required = [c.strip().lower() for c in required]
    return [c for c in required if c not in col_map]
